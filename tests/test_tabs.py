"""Which tab of the workbook holds this week's board.

The regression: `wb[wb.sheetnames[0]]`. The operator's week 2 workbook is
['01', 'standing', 'cal', 'Week02'] -- tab one is last week's finished results. On that
specific file the old code halted, because his results tab carries no header row, so a
run on 2026-09-20 would have produced no board at all. It is the OTHER shape that these
tests mostly exist for: a results tab that does keep its header parses perfectly well and
yields a complete, confident board for a week that has already been played.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from madden.sheet import SheetFault, count_games, find_header, parse_sheet, select_tab


def board(ws, rows, scored=False):
    """Write a board onto a worksheet. `scored` adds a result column, as a played week has."""
    ws["C1"] = "2026-2027 NFL Picks"
    ws["E6"], ws["J6"], ws["O6"] = "Favorite", "Spread", "Underdog"
    r = 8
    for fav, spread, dog in rows:
        ws.cell(row=r, column=5, value=fav)
        ws.cell(row=r, column=10, value=spread)
        ws.cell(row=r, column=15, value=dog)
        if scored:
            ws.cell(row=r, column=20, value="24-20")
        r += 2


def results_tab_without_header(ws, rows):
    """The operator's real results layout: numbered rows, no labels anywhere."""
    r = 5
    for i, (fav, spread, dog) in enumerate(rows, start=1):
        ws.cell(row=r, column=1, value=i)
        ws.cell(row=r, column=2, value=fav)
        ws.cell(row=r, column=3, value=spread)
        ws.cell(row=r, column=4, value=dog)
        ws.cell(row=r, column=5, value="24-20")
        r += 1


THIS_WEEK = [("BUFFALO BILLS", 4.5, "Detroit Lions"),
             ("Green Bay Packers", 1.5, "MINNESOTA VIKINGS"),
             ("LOS ANGELES RAMS", 6.5, "New York Giants")]
LAST_WEEK = [("SEATTLE SEAHAWKS", 3.5, "New England Patriots"),
             ("JACKSONVILLE JAGUARS", 7.5, "Cleveland Browns")]


# ---------------------------------------------------------------- case 1: one tab passes

def test_the_board_is_found_past_a_results_tab_that_has_no_header(tmp_path):
    """The real week 2 shape: ['01', 'standing', 'cal', 'Week02']."""
    wb = Workbook()
    results_tab_without_header(wb.active, LAST_WEEK)
    wb.active.title = "01"
    wb.create_sheet("standing")["A1"] = "Season Results by Week"
    wb.create_sheet("cal")["A1"] = "Actual Record"
    board(wb.create_sheet("Week02"), THIS_WEEK)
    p = tmp_path / "w02.xlsx"
    wb.save(p)

    seen = []
    games = parse_sheet(str(p), expected_games=3, announce=seen.append)
    assert [g.favorite for g in games] == ["BUF", "GB", "LAR"]
    assert "'Week02'" in seen[0], "the tab chosen is announced"


def test_a_single_tab_workbook_still_works(tmp_path):
    """Week 1's file is one tab called Week01. The announcement stays quiet about others."""
    wb = Workbook()
    wb.active.title = "Week01"
    board(wb.active, THIS_WEEK)
    p = tmp_path / "w01.xlsx"
    wb.save(p)
    seen = []
    assert len(parse_sheet(str(p), announce=seen.append)) == 3
    assert seen == ["tab:   'Week01'"]


# ------------------------------------------------- case 2: several pass, most games wins

def test_when_several_tabs_parse_the_one_with_more_games_wins(tmp_path):
    """The dangerous shape: a results tab that KEPT its header and parses cleanly.

    Nothing about the tab order, the names or the scores decides this. Only the count.
    """
    wb = Workbook()
    wb.active.title = "01"
    board(wb.active, LAST_WEEK, scored=True)          # 2 games, first in the book
    board(wb.create_sheet("Week02"), THIS_WEEK)       # 3 games
    p = tmp_path / "both.xlsx"
    wb.save(p)

    seen = []
    games = parse_sheet(str(p), announce=seen.append)
    assert [g.favorite for g in games] == ["BUF", "GB", "LAR"]
    assert "'Week02'" in seen[0]


def test_order_does_not_decide_it(tmp_path):
    """Same workbook with the tabs reversed picks the same tab."""
    wb = Workbook()
    wb.active.title = "Week02"
    board(wb.active, THIS_WEEK)
    board(wb.create_sheet("01"), LAST_WEEK, scored=True)
    p = tmp_path / "reversed.xlsx"
    wb.save(p)
    assert len(parse_sheet(str(p))) == 3


def test_two_tabs_parsing_equally_well_is_a_halt(tmp_path):
    """Ambiguous is not resolved by a coin toss. It names both and stops."""
    wb = Workbook()
    wb.active.title = "Week01"
    board(wb.active, THIS_WEEK)
    board(wb.create_sheet("Week02"), THIS_WEEK)
    p = tmp_path / "tie.xlsx"
    wb.save(p)
    with pytest.raises(SheetFault) as exc:
        parse_sheet(str(p))
    assert "ambiguous" in str(exc.value)
    assert "'Week01'" in str(exc.value) and "'Week02'" in str(exc.value)


# ------------------------------------------------------ case 3: none pass, halt and name

def test_no_tab_with_a_header_halts_and_names_every_tab(tmp_path):
    wb = Workbook()
    wb.active.title = "01"
    results_tab_without_header(wb.active, LAST_WEEK)
    wb.create_sheet("standing")["A1"] = "Season Results by Week"
    wb.create_sheet("cal")["A1"] = "Actual Record"
    p = tmp_path / "none.xlsx"
    wb.save(p)

    with pytest.raises(SheetFault) as exc:
        parse_sheet(str(p))
    message = str(exc.value)
    assert "no tab" in message
    for name in ("'01'", "'standing'", "'cal'"):
        assert name in message, f"{name} must be named in the refusal"
    assert "Examined 3" in message


def test_a_header_with_nothing_under_it_halts_and_says_so(tmp_path):
    wb = Workbook()
    wb.active.title = "Week02"
    wb.active["E6"], wb.active["J6"], wb.active["O6"] = "Favorite", "Spread", "Underdog"
    p = tmp_path / "empty.xlsx"
    wb.save(p)
    with pytest.raises(SheetFault) as exc:
        parse_sheet(str(p))
    assert "no game parses" in str(exc.value)


# ------------------------------------------------------------------------- the two pieces

def test_find_header_locates_labels_not_a_row_number():
    rows = [[None], ["Favorite", "Spread", "Underdog"], ["x"]]
    h = find_header(rows)
    assert (h.row, h.favorite, h.spread, h.underdog) == (1, 0, 1, 2)
    assert find_header([["Favorite", "Spread"]]) is None, "all three labels or nothing"


def test_count_games_never_raises_on_a_tab_it_is_only_scoring():
    """A bad row on a tab about to be discarded is not a sheet fault."""
    rows = [["Favorite", "Spread", "Underdog"],
            ["BUFFALO BILLS", 4.5, "Detroit Lions"],      # good
            ["Not A Team", 3.5, "Also Not A Team"],       # unresolvable
            ["BUFFALO BILLS", -1.0, "Detroit Lions"],     # negative spread
            ["BUFFALO BILLS", 4.5, "DETROIT LIONS"]]      # both caps: no home team
    assert count_games(rows, find_header(rows)) == 1


def test_select_tab_returns_the_rows_it_chose(tmp_path):
    wb = Workbook()
    wb.active.title = "01"
    board(wb.active, LAST_WEEK, scored=True)
    board(wb.create_sheet("Week02"), THIS_WEEK)
    p = tmp_path / "s.xlsx"
    wb.save(p)
    from openpyxl import load_workbook
    name, rows, header = select_tab(load_workbook(p, data_only=True, read_only=True))
    assert name == "Week02" and count_games(rows, header) == 3
