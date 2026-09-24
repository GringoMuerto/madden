"""Tests that pin the things which silently invert the whole system.

The sign convention is the dangerous one. Situational points are added to the HOME team,
so a negative adjustment favours the visitor. A previous reconstruction had this rule at
-3.3 with an invented 70F companion rule, and it flipped picks in six games.
"""

import http.client
import sys
from datetime import datetime, timezone
from pathlib import Path

import pytest
import yaml

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from madden.core import band_for, crosses_key_number, make_pick, situational_adjustments
from madden.sheet import Game, SheetFault, parse_sheet

PARAMS = yaml.safe_load((Path(__file__).resolve().parents[1] / "params.yaml").read_text())


def game(fav, dog, spread, home, day="Sunday", neutral=False):
    g = Game(favorite=fav, underdog=dog, spread=spread, day=day, nominal_home=home, row=0)
    g.neutral_site = neutral
    return g


def test_spec_values_are_exact():
    s = PARAMS["situational"]
    assert s["dome_visitor_outdoors_75f"]["points"] == -2.8
    assert s["home_underdog_7plus"]["points"] == 2.3
    assert s["divisional"]["points"] == -0.8
    assert s["global_home_baseline"]["points"] == -0.3
    assert "outdoor_70f" not in s, "there is no 70F rule in the spec"


def test_global_baseline_applies_to_every_game():
    g = game("BAL", "IND", 3.5, home="IND")          # no other rule fires
    adj, _ = situational_adjustments(g, PARAMS)
    assert [a.rule for a in adj] == ["global_home_baseline"]


def test_baseline_alone_takes_the_road_team():
    g = game("BAL", "IND", 3.5, home="IND")
    p = make_pick(g, market_home_line=-3.5, params=PARAMS)
    assert p.edge == -0.3 and p.side == "BAL" and p.band == "coin flip"


def test_dome_visitor_outdoors_favours_the_visitor():
    # Atlanta is dome-based, Pittsburgh is outdoors at 79F.
    g = game("PIT", "ATL", 2.5, home="PIT")
    adj, _ = situational_adjustments(g, PARAMS, temp_f=79)
    assert {a.rule for a in adj} == {"global_home_baseline", "dome_visitor_outdoors_75f"}
    p = make_pick(g, market_home_line=3.5, params=PARAMS, temp_f=79)
    assert p.adjustment_total == -3.1
    assert p.side == "ATL" and p.band == "medium"


def test_rule_does_not_fire_below_75f():
    g = game("PIT", "ATL", 2.5, home="PIT")
    adj, _ = situational_adjustments(g, PARAMS, temp_f=74)
    assert all(a.rule != "dome_visitor_outdoors_75f" for a in adj)


def test_rule_does_not_fire_indoors():
    # New Orleans is dome-based but Detroit plays under a fixed roof.
    g = game("DET", "NO", 6.5, home="DET")
    adj, _ = situational_adjustments(g, PARAMS, temp_f=90)
    assert all(a.rule != "dome_visitor_outdoors_75f" for a in adj)


def test_divisional_and_dome_visitor_stack():
    # Dallas is dome-based, MetLife is outdoors at 77F, and it is an NFC East game.
    g = game("DAL", "NYG", 2.5, home="NYG")
    p = make_pick(g, market_home_line=-2.5, params=PARAMS, temp_f=77)
    assert p.adjustment_total == -3.9
    assert p.madden_number == -6.4 and p.edge == -3.9
    assert p.side == "DAL" and p.band == "high"


def test_home_underdog_of_seven_plus():
    g = game("KC", "CAR", 8.5, home="CAR")
    adj, _ = situational_adjustments(g, PARAMS, temp_f=60)
    assert any(a.rule == "home_underdog_7plus" and a.points == 2.3 for a in adj)


def test_neutral_site_voids_home_rules():
    g = game("LAR", "SF", 3.5, home="LAR", day="Thursday", neutral=True)
    adj, warn = situational_adjustments(g, PARAMS, temp_f=60)
    assert [a.rule for a in adj] == ["divisional"]
    assert any("neutral site" in w for w in warn)


def test_key_number_upgrades_to_high():
    # Chargers 10.5 on the sheet, 9.5 in the market: the edge steps across 10.
    g = game("LAC", "ARI", 10.5, home="LAC")
    p = make_pick(g, market_home_line=9.5, params=PARAMS, temp_f=70)
    assert p.edge == -1.3
    assert p.band == "high" and p.side == "ARI"


def test_crossing_detection_handles_home_dogs():
    assert crosses_key_number(-6.4, -2.5, [3, 7, 10, 14]) == 3
    assert crosses_key_number(3.2, 3.5, [3, 7, 10, 14]) is None


def test_missing_line_is_never_fabricated():
    g = game("BAL", "IND", 3.5, home="IND")
    p = make_pick(g, market_home_line=None, params=PARAMS)
    assert p.side is None and p.blind and p.madden_number is None


def test_a_drift_past_the_threshold_withholds_the_pick():
    """The week 1 SF at LAR row, which is why this check exists.

    Sheet +3.5, recorded market -19.0, a drift of 22.5. The old build priced it, banded it
    high on the size of the error and printed "line moved 22.5 toward SF" as the driver.
    Nothing warned. The pick won, which is why it went unnoticed.
    """
    # LAR was the sheet's favorite and its nominal home, so the sheet home line is +3.5.
    g = game("LAR", "SF", 3.5, home="LAR", day="Thursday")
    p = make_pick(g, market_home_line=-19.0, params=PARAMS)
    assert p.side is None, "a line this far out must not produce a pick"
    assert p.band == "blind" and p.madden_number is None and p.edge is None
    # The rejected number is kept: it is the finding, not noise to be blanked.
    assert p.market_home_line == -19.0
    assert any("suspect" in w for w in p.warnings)
    assert any("22.5" in w for w in p.warnings), "the warning names the drift it saw"


def test_a_drift_inside_the_threshold_still_prices_normally():
    """The guard is set to catch a broken feed row, never to second-guess a real move."""
    limit = PARAMS["odds_api"]["max_drift_points"]
    g = game("PIT", "ATL", 2.5, home="PIT")
    p = make_pick(g, market_home_line=2.5 + limit - 0.5, params=PARAMS, temp_f=70)
    assert p.side is not None and p.madden_number is not None
    assert not any("suspect" in w for w in p.warnings)


def test_the_drift_threshold_is_required_not_defaulted():
    """A missing key fails loudly rather than disabling the check.

    The spec's own recorded failure is a rule that matched nothing and read exactly like
    protection for days. A .get() with a fallback here would rebuild that.
    """
    stripped = {k: v for k, v in PARAMS.items() if k != "odds_api"}
    g = game("BAL", "IND", 3.5, home="IND")
    with pytest.raises(KeyError):
        make_pick(g, market_home_line=3.0, params=stripped)


def test_band_thresholds():
    assert band_for(3.0, 0, 3.0, PARAMS)[0] == "high"
    assert band_for(1.5, 20.0, 18.5, PARAMS)[0] == "medium"
    assert band_for(0.4, 20.0, 19.6, PARAMS)[0] == "coin flip"


def _write_sheet(path, rows):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws["C1"] = "2026-2027 NFL Picks"
    ws["C5"] = "HOME TEAM IN CAPITAL LETTERS."
    ws["E6"], ws["G6"], ws["J6"], ws["M6"], ws["O6"] = (
        "Favorite", "W-L", "Spread", "W-L", "Underdog")
    r = 8
    for day, fav, spread, dog in rows:
        if day:
            ws.cell(row=r, column=10, value=f"{day} Games")
            r += 2
        ws.cell(row=r, column=5, value=fav)
        ws.cell(row=r, column=10, value=spread)
        ws.cell(row=r, column=15, value=dog)
        r += 2
    wb.save(path)


def test_parser_reads_layout_and_home_capitalisation(tmp_path):
    p = tmp_path / "sheet.xlsx"
    _write_sheet(p, [
        ("Thursday", "LOS ANGELES RAMS", 3.5, "San Francisco 49ers"),
        ("Sunday", "Baltimore Ravens", 3.5, "INDIANAPOLIS COLTS"),
        (None, "KANSAS CITY CHIEFS", 2.5, "Denver Broncos"),
    ])
    games = parse_sheet(str(p), expected_games=3)
    assert [g.day for g in games] == ["Thursday", "Sunday", "Sunday"]
    assert games[0].home == "LAR" and games[0].sheet_home_line == 3.5
    assert games[1].home == "IND" and games[1].sheet_home_line == -3.5
    assert games[2].favorite == "KC"


def test_game_count_mismatch_is_a_hard_stop(tmp_path):
    p = tmp_path / "sheet.xlsx"
    _write_sheet(p, [("Sunday", "KANSAS CITY CHIEFS", 2.5, "Denver Broncos")])
    with pytest.raises(SheetFault):
        parse_sheet(str(p), expected_games=16)


def test_neutral_site_declaration_flows_through_the_parser(tmp_path):
    p = tmp_path / "sheet.xlsx"
    _write_sheet(p, [("Thursday", "LOS ANGELES RAMS", 3.5, "San Francisco 49ers")])
    games = parse_sheet(str(p), neutral_sites=[("LAR", "SF")])
    assert games[0].neutral_site


def _event(home, away, home_spread, kickoff):
    """One odds-api event with a single book quoting the home team at `home_spread`."""
    return {
        "home_team": home, "away_team": away, "commence_time": kickoff,
        "bookmakers": [{"last_update": kickoff, "markets": [{"key": "spreads", "outcomes": [
            {"name": home, "point": home_spread}, {"name": away, "point": -home_spread}]}]}],
    }


# Market tests run at a fixed moment, never the wall clock: soonest() drops a game once it
# has kicked off, so a test pinned to real kickoffs breaks the day those games are played.
WEEK1_TUESDAY = datetime(2026, 9, 8, 12, 0, tzinfo=timezone.utc)


def test_divisional_rematch_does_not_overwrite_this_weeks_line():
    # Week 1 at the Giants, NYG -2.5. The week 17 rematch in Dallas, DAL -4.5, comes back
    # in the same payload. The run must price week 1, oriented to the sheet's home team.
    # A rematch months away is a different game, not an ambiguity, so it must NOT warn:
    # warning on the normal case trains the reader to skip run health (commit 0fcb959).
    from madden.market import parse_odds_payload, soonest
    payload = [
        _event("Dallas Cowboys", "New York Giants", -4.5, "2026-12-27T18:00:00Z"),
        _event("New York Giants", "Dallas Cowboys", -2.5, "2026-09-13T17:00:00Z"),
    ]
    lines = parse_odds_payload(payload, PARAMS)
    ml, warnings = soonest(lines, "NYG", "DAL", now=WEEK1_TUESDAY)
    assert ml.commence_time.startswith("2026-09-13")
    assert ml.line_for("NYG") == 2.5
    assert not any("meetings" in w for w in warnings)


def test_two_meetings_inside_the_window_do_warn():
    from madden.market import parse_odds_payload, soonest
    payload = [
        _event("New York Giants", "Dallas Cowboys", -2.5, "2026-09-13T17:00:00Z"),
        _event("Dallas Cowboys", "New York Giants", -4.5, "2026-09-17T00:15:00Z"),
    ]
    ml, warnings = soonest(parse_odds_payload(payload, PARAMS), "NYG", "DAL",
                           now=WEEK1_TUESDAY)
    assert ml.commence_time.startswith("2026-09-13")
    assert any("2 meetings" in w for w in warnings)


def test_a_game_already_kicked_off_has_no_line():
    from madden.market import parse_odds_payload, soonest
    payload = [_event("New York Giants", "Dallas Cowboys", -2.5, "2026-09-13T17:00:00Z")]
    after = datetime(2026, 9, 13, 18, 0, tzinfo=timezone.utc)
    ml, warnings = soonest(parse_odds_payload(payload, PARAMS), "NYG", "DAL", now=after)
    assert ml is None and any("already kicked off" in w for w in warnings)


def test_line_is_reoriented_when_the_feed_disagrees_about_home():
    from madden.market import parse_odds_payload, soonest
    payload = [_event("San Francisco 49ers", "Los Angeles Rams", -1.5, "2026-09-11T10:35:00Z")]
    ml, warnings = soonest(parse_odds_payload(payload, PARAMS), "LAR", "SF",
                           now=WEEK1_TUESDAY)
    assert ml.line_for("LAR") == -1.5
    assert any("re-oriented" in w for w in warnings)


# Quarterback exposure, against a canned nflverse. Names only: nothing here may carry a
# number, and a team with no data is UNKNOWN, never clean.

INJ_HEADER = ("season,season_type,game_type,team,week,gsis_id,position,full_name,first_name,"
              "last_name,report_primary_injury,report_status,practice_primary_injury,"
              "practice_secondary_injury,practice_status")


def _nflverse(rows, depth=(), broken=()):
    """A fake GET. rows: (team, week, gsis, pos, name, report_status, practice_status, injury)."""
    import gzip
    import json as _json
    from madden.injuries import RELEASE_API, RELEASES
    lines = [INJ_HEADER] + [
        f"2026,REG,REG,{t},{w},{g},{pos},{n},,,{inj},{rs},,,{ps}"
        for t, w, g, pos, n, rs, ps, inj in rows]
    dlines = ["dt,team,player_name,espn_id,gsis_id,pos_grp_id,pos_grp,pos_id,pos_name,"
              "pos_abb,pos_slot,pos_rank"] + [
        f"2026-09-10T12:01:46Z,{t},{n},,{g},21,3WR 1TE,8,Quarterback,QB,8,{rank}"
        for t, g, n, rank in depth]
    table = {
        f"{RELEASES}/injuries/injuries_2026.csv": "\n".join(lines).encode(),
        f"{RELEASE_API}/injuries": _json.dumps({"assets": [
            {"name": "injuries_2026.csv", "updated_at": "2026-09-13T12:03:24Z"}]}).encode(),
        f"{RELEASES}/depth_charts/depth_charts_2026.csv.gz":
            gzip.compress("\n".join(dlines).encode()),
    }

    def get(url):
        if any(b in url for b in broken) or url not in table:
            raise RuntimeError(f"HTTP Error 404 for {url}")
        return table[url]
    return get


ATL_DEPTH = [("ATL", "g-tua", "Tua Tagovailoa", 1), ("ATL", "g-penix", "Michael Penix Jr.", 2)]


def test_exposure_names_an_unresolved_quarterback_with_rank_and_no_points():
    from madden import injuries
    rep, _ = injuries.fetch(2026, 1, get=_nflverse([
        ("ATL", 1, "g-tua", "QB", "Tua Tagovailoa", "Questionable",
         "Limited Participation in Practice", "Oblique"),
        ("PIT", 1, "g-x", "CB", "Donte Kent", "Out", "", "Knee"),
    ], depth=ATL_DEPTH))
    names, unknown = injuries.exposure(rep, "PIT", "ATL")
    assert names == ["ATL QB Tua Tagovailoa (QB1), Questionable (Oblique)"]
    assert unknown == []
    assert rep.built == "2026-09-13T12:03:24Z"


def test_out_and_full_participation_are_resolved():
    from madden import injuries
    rep, _ = injuries.fetch(2026, 1, get=_nflverse([
        ("ATL", 1, "g-penix", "QB", "Michael Penix Jr.", "Out", "", "Knee"),
        ("WAS", 1, "g-mar", "QB", "Marcus Mariota", "", "Full Participation in Practice", "Knee"),
    ], depth=ATL_DEPTH))
    assert injuries.exposure(rep, "PIT", "ATL") == ([], ["PIT"])
    assert injuries.exposure(rep, "PHI", "WAS") == ([], ["PHI"])


def test_no_game_status_counts_only_until_the_final_report():
    from madden import injuries
    dnp = ("CHI", 1, "g-bag", "QB", "Tyson Bagent", "", "Did Not Participate In Practice", "Back")
    rep, _ = injuries.fetch(2026, 1, get=_nflverse([dnp, ("CAR", 1, "g-c", "OT", "X", "", "", "")]))
    names, _ = injuries.exposure(rep, "CAR", "CHI")
    assert names == ["CHI QB Tyson Bagent (not on the depth chart), no game status yet, "
                     "did not practice (Back)"]
    # Once any Chicago player carries a game status, the final report is out and a blank
    # status means no designation.
    final = ("CHI", 1, "g-g", "CB", "Kyler Gordon", "Out", "", "Hamstring")
    rep, _ = injuries.fetch(2026, 1, get=_nflverse([dnp, final, ("CAR", 1, "g-c", "OT", "X", "", "", "")]))
    assert injuries.exposure(rep, "CAR", "CHI") == ([], [])


def test_team_missing_from_the_report_is_unknown_not_clean():
    from madden import injuries
    rep, _ = injuries.fetch(2026, 1, get=_nflverse([("DEN", 2, "g-n", "QB", "Bo Nix", "", "", "")]))
    assert injuries.exposure(rep, "KC", "DEN") == ([], ["DEN", "KC"])


def test_report_fetch_failure_is_unknown_for_every_game():
    from madden import injuries
    rep, _ = injuries.fetch(2026, 1, get=_nflverse([], broken=("injuries_2026.csv",)))
    assert rep.error and "404" in rep.error
    assert injuries.exposure(rep, "PIT", "ATL") == ([], ["ATL", "PIT"])
    assert injuries.header(rep).startswith("INJURY REPORT FETCH FAILED")


def test_depth_chart_failure_names_without_rank_and_warns():
    from madden import injuries
    rep, warns = injuries.fetch(2026, 1, get=_nflverse([
        ("ATL", 1, "g-tua", "QB", "Tua Tagovailoa", "Doubtful", "", "Oblique"),
    ], broken=("depth_charts",)))
    assert injuries.exposure(rep, "PIT", "ATL")[0] == ["ATL QB Tua Tagovailoa, Doubtful (Oblique)"]
    assert any("depth chart unavailable" in w for w in warns)


def test_nflverse_rams_are_lar():
    from madden import injuries
    rep, _ = injuries.fetch(2026, 1, get=_nflverse([
        ("LA", 1, "g-s", "QB", "Matthew Stafford", "Questionable", "", "Back")]))
    assert injuries.exposure(rep, "LAR", "SF")[0] == [
        "LAR QB Matthew Stafford (not on the depth chart), Questionable (Back)"]


def test_forecast_is_not_written_to_disk_without_the_cache_flag(tmp_path, monkeypatch):
    from madden import weather
    monkeypatch.setattr(weather, "CACHE", tmp_path / ".cache")
    monkeypatch.setattr(weather, "_fetch", lambda lat, lon, timeout: {"hourly": {
        "time": ["2026-09-13T17:00"], "temperature_2m": [80.4], "wind_speed_10m": [5.0]}})
    kickoff = datetime(2026, 9, 13, 17, 0, tzinfo=timezone.utc)
    assert weather.forecast("PIT", kickoff) == (80.4, 5.0)
    assert not (tmp_path / ".cache").exists()


# Under the sandbox proxy a request sent with "Connection: close" can come back with the
# tail of its body missing (madden/net.py). Each fetcher's request is captured at the
# point http.client would send it, so nothing reaches the network.

class _Captured(Exception):
    pass


def _fetchers():
    from madden import injuries, market, weather
    return {
        "market": lambda: market.fetch_lines(PARAMS, api_key="test-key"),
        "injuries": lambda: injuries._http(f"{injuries.RELEASES}/injuries/injuries_2026.csv"),
        "weather": lambda: weather._fetch(40.447, -80.016, timeout=5),
    }


@pytest.mark.parametrize("name", ["market", "injuries", "weather"])
def test_no_fetcher_sends_connection_close(name, monkeypatch):
    import http.client
    sent = []

    def capture(self, message_body=None, encode_chunked=False):
        sent.append(b"\r\n".join(self._buffer).decode("latin-1"))
        raise _Captured

    monkeypatch.setattr(http.client.HTTPConnection, "_send_output", capture)
    with pytest.raises(_Captured):
        _fetchers()[name]()
    # _buffer is http.client's private header list; this line stops a silent pass if
    # a future Python stops building the request there.
    assert sent and sent[0].startswith("GET ")
    headers = [line.lower() for line in sent[0].split("\r\n")[1:]]
    assert not any(h.startswith("connection:") for h in headers), sent[0]


@pytest.mark.parametrize("exc", [
    http.client.IncompleteRead(b'{"hourly":', 5000),
    http.client.BadStatusLine(""),
], ids=["incomplete-read", "bad-status-line"])
def test_a_broken_forecast_degrades_rather_than_stopping_the_run(exc, monkeypatch):
    """A cut-off or malformed forecast is a data fault: warn and carry on, never halt.

    These are http.client.HTTPException, which is neither OSError nor ValueError, so
    they walked straight out of forecast() and would have ended the run.
    """
    from madden import weather

    def fail(lat, lon, timeout):
        raise exc

    monkeypatch.setattr(weather, "_fetch", fail)
    kickoff = datetime(2026, 9, 13, 17, 0, tzinfo=timezone.utc)
    temps, winds, warnings = weather.forecast_many([("PIT", kickoff)])
    assert temps == {} and winds == {}
    assert any("no forecast for PIT" in w for w in warnings)


def _handback(side, warnings=("flagged",)):
    """A blind pick with a lean (side set) or with no pick at all (side None)."""
    from madden.core import Pick
    g = game("KC", "DEN", 2.5, home="KC") if side else game("BAL", "IND", 3.5, home="IND")
    return Pick(game=g, market_home_line=-2.5 if side else None,
                sheet_home_line=g.sheet_home_line, adjustments=[], madden_number=None,
                edge=None, side=side, band="blind", warnings=list(warnings), blind=True)


def test_handbacks_claim_a_lean_only_for_games_that_have_one():
    from madden.run import handback_lines
    out = handback_lines([_handback("KC"), _handback(None)])
    lean_header = next(i for i, l in enumerate(out) if l.startswith("HANDBACKS"))
    no_pick_header = next(i for i, l in enumerate(out) if l.startswith("NO PICK"))
    leaned = next(i for i, l in enumerate(out) if "lean KC" in l)
    unpicked = next(i for i, l in enumerate(out) if "BAL at IND" in l)
    assert lean_header < leaned < no_pick_header < unpicked
    assert "lean" not in out[unpicked]


def test_no_lean_is_claimed_when_no_game_carries_one():
    from madden.run import handback_lines
    out = handback_lines([_handback(None)])
    assert out and not any("carries a lean" in line for line in out)
    assert any("revert to the favorite" in line for line in out)


# The spreads feed drops a game once it has been played, so a finished game arrives
# looking exactly like one the feed never listed: no line. The scores endpoint separates
# them, and a failed fetch must never be reported as either.

SCORES_PAYLOAD = [
    {"home_team": "Seattle Seahawks", "away_team": "New England Patriots",
     "commence_time": "2026-09-10T00:23:22Z", "completed": True,
     "scores": [{"name": "Seattle Seahawks", "score": "13"},
                {"name": "New England Patriots", "score": "10"}]},
    {"home_team": "Kansas City Chiefs", "away_team": "Denver Broncos",
     "commence_time": "2026-09-13T17:00:00Z", "completed": False, "scores": None},
]


def test_scores_payload_keys_on_the_pair_and_keeps_the_final():
    from madden.market import parse_scores_payload
    scores = parse_scores_payload(SCORES_PAYLOAD)
    played = scores[frozenset(("SEA", "NE"))]
    assert played.completed and "Seattle Seahawks 13" in played.text
    assert not scores[frozenset(("KC", "DEN"))].completed


def test_a_played_game_is_distinguishable_from_a_failed_fetch():
    from madden.market import parse_scores_payload
    from madden.run import no_line_reason
    scores = parse_scores_payload(SCORES_PAYLOAD)
    played = game("SEA", "NE", 3.5, home="SEA")
    never_listed = game("BAL", "IND", 3.5, home="IND")

    played_label, played_text = no_line_reason(played, scores, fetch_failed=False)
    absent_label, _ = no_line_reason(never_listed, scores, fetch_failed=False)
    failed_label, failed_text = no_line_reason(played, scores, fetch_failed=True)

    assert played_label == "already played" and "Seattle Seahawks 13" in played_text
    assert absent_label == "not in the feed"
    # Same game, but with the fetch broken nothing may be claimed about it.
    assert failed_label == "fetch failed" and "line fetch failed" in failed_text
    assert len({played_label, absent_label, failed_label}) == 3


# Which week the sheet is. Derived from the games, never from the filename (the operator
# numbers his files from zero) and never from the clock. An unresolved week halts, because
# the injury report is keyed on it and a wrong week goes dark without a warning.

SCHEDULE_HEADER = "game_id,season,game_type,week,gameday,weekday,gametime,away_team,home_team"

WEEK1 = [("NE", "SEA"), ("SF", "LA"), ("CHI", "CAR"), ("TB", "CIN"), ("NO", "DET"),
         ("BUF", "HOU"), ("BAL", "IND"), ("CLE", "JAX"), ("ATL", "PIT"), ("NYJ", "TEN"),
         ("ARI", "LAC"), ("MIA", "LV"), ("GB", "MIN"), ("WAS", "PHI"), ("DAL", "NYG"),
         ("DEN", "KC")]
# Sixteen pairings that share none of week 1's. Pairs are matched unordered, so a week
# built by reversing week 1's home and away would be the same sixteen pairs and would tie.
WEEK2 = [("DET", "BUF"), ("CAR", "ATL"), ("NO", "BAL"), ("MIN", "CHI"), ("SEA", "PIT"),
         ("LA", "TEN"), ("CIN", "CLE"), ("HOU", "JAX"), ("IND", "KC"), ("NYJ", "MIA"),
         ("LAC", "LV"), ("GB", "PHI"), ("NYG", "WAS"), ("DAL", "ARI"), ("SF", "DEN"),
         ("NE", "TB")]


def _schedule(weeks=None, broken=False):
    """A fake GET returning an nflverse games.csv. weeks: {(season, week): [(away, home)]}."""
    weeks = {(2026, 1): WEEK1, (2026, 2): WEEK2} if weeks is None else weeks
    lines = [SCHEDULE_HEADER]
    for (season, wk), pairs in weeks.items():
        for i, (away, home) in enumerate(pairs):
            lines.append(f"{season}_{wk}_{away}_{home},{season},REG,{wk},"
                         f"2026-09-{12 + wk:02d},Sunday,13:00,{away},{home}")

    def get(url):
        if broken:
            raise RuntimeError("HTTP Error 404")
        return "\n".join(lines).encode()
    return get


def _sheet_games(pairs):
    """The same pairings as the sheet would carry them.

    The lists above are spelled the way nflverse spells them, so the Rams are LA. The
    sheet says "Los Angeles Rams" and the parser resolves that to LAR, so the sheet side
    of every match is exercised in our spelling, not the feed's.
    """
    ours = {"LA": "LAR"}
    return [game(ours.get(away, away), ours.get(home, home), 3.5, home=ours.get(home, home))
            for away, home in pairs]


def test_week_comes_from_the_games_not_the_filename():
    from madden import schedule
    # The real sheet: filename NFL2026w0.xlsx, tab "Week01", and these sixteen games.
    res = schedule.resolve(_sheet_games(WEEK1), get=_schedule())
    assert (res.season, res.week) == (2026, 1)
    assert res.matched == res.total == 16
    assert res.source == "nflverse schedule"


def test_the_operators_zero_indexed_filename_is_not_warned_about():
    from madden import schedule
    res = schedule.resolve(_sheet_games(WEEK1), get=_schedule())
    # w0 for NFL week 1 is his habit, so it is expected, not a warning.
    assert schedule.filename_cross_check(res, 0) == []
    assert schedule.filename_cross_check(res, 1) == []
    assert schedule.filename_cross_check(res, None) == []
    # Anything else means the wrong file may have been picked up.
    off = schedule.filename_cross_check(res, 7)
    assert off and "NFL week 1" in off[0] and "numbers his files from zero" in off[0]


def test_the_rams_are_matched_through_the_nflverse_spelling():
    from madden import schedule
    # nflverse writes the Rams as LA; the sheet writes LAR. SF at LAR must still match.
    res = schedule.resolve(_sheet_games([("SF", "LAR")]), get=_schedule())
    assert (res.season, res.week) == (2026, 1)


def test_a_sheet_matching_no_week_is_a_hard_stop():
    from madden import schedule
    nonsense = _sheet_games([("KC", "DEN"), ("SF", "SEA"), ("NYG", "DAL"), ("MIA", "BUF")])
    with pytest.raises(SheetFault) as exc:
        schedule.resolve(nonsense, get=_schedule(weeks={(2026, 1): WEEK1}))
    assert "match no week" in str(exc.value)


def test_an_unreachable_schedule_halts_rather_than_going_dark():
    from madden import schedule
    with pytest.raises(SheetFault) as exc:
        schedule.resolve(_sheet_games(WEEK1), get=_schedule(broken=True))
    # The halt has to name the escape hatch, or a broken nflverse ends the week.
    assert "dark exposure layer" in str(exc.value)
    assert "--week" in str(exc.value)


def test_an_ambiguous_week_halts_rather_than_picking_one():
    from madden import schedule
    twice = {(2026, 1): WEEK1, (2026, 9): WEEK1}
    with pytest.raises(SheetFault) as exc:
        schedule.resolve(_sheet_games(WEEK1), get=_schedule(weeks=twice))
    assert "ambiguous" in str(exc.value)


def test_a_stale_week_file_is_reported_rather_than_obeyed_silently():
    from madden import schedule
    games = _sheet_games(WEEK1)
    assert schedule.confirm(2026, 1, games, get=_schedule()) == []
    stale = schedule.confirm(2026, 2, games, get=_schedule())
    assert stale and "week file wins" in stale[0] and "leftover" in stale[0]
    # The override still has to work when nflverse is unreachable: warn, never halt.
    offline = schedule.confirm(2026, 2, games, get=_schedule(broken=True))
    assert offline and "could not be confirmed" in offline[0]


def test_a_clean_fetch_of_an_empty_week_is_not_mistaken_for_a_healthy_one():
    from madden import injuries
    healthy, _ = injuries.fetch(2026, 1, get=_nflverse([
        ("ATL", 1, "g-tua", "QB", "Tua Tagovailoa", "Questionable", "", "Oblique")]))
    assert not healthy.empty
    # The quiet failure: right URL, wrong week. HTTP 200, no rows, every team UNKNOWN.
    wrong_week, _ = injuries.fetch(2026, 1, get=_nflverse([
        ("ATL", 5, "g-tua", "QB", "Tua Tagovailoa", "Questionable", "", "Oblique")]))
    assert wrong_week.empty and not wrong_week.error
    assert injuries.exposure(wrong_week, "PIT", "ATL") == ([], ["ATL", "PIT"])
    # A failed fetch is a different state and must stay distinguishable from an empty one.
    broken, _ = injuries.fetch(2026, 1, get=_nflverse([], broken=("injuries_2026.csv",)))
    assert broken.error and not broken.empty


# The power rating. Not built, by design, and a run says nothing about a design decision.
# It speaks only when the config claims a layer the engine does not have.

def test_the_run_is_silent_about_the_rating_not_being_built():
    from madden.run import rating_module
    assert rating_module() is None, "no rating is built; nothing should claim otherwise"
    assert PARAMS["power_rating"]["enabled"] is False
    # Both together are the normal state, and the normal state produces no warning.
    assert not (PARAMS["power_rating"]["enabled"] and rating_module() is None)


def test_the_model_term_is_zero_whatever_the_flag_says():
    # Flipping the flag must not move a number: run.py passes no model_term and
    # make_pick defaults it to 0.0. The flag is intent; the module is the fact.
    import copy
    p = copy.deepcopy(PARAMS)
    g = game("DEN", "KC", 2.5, home="KC")
    before = make_pick(g, 2.5, p).madden_number
    p["power_rating"]["enabled"] = True
    after = make_pick(g, 2.5, p).madden_number
    assert before == after
